"""Exercise document functions without importing the app's production database."""
import ast
import asyncio
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import threading
import time
from typing import List, Optional
import unittest
from unittest.mock import patch
import zipfile

from fastapi import FastAPI, Depends, File, Form, Header, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from PIL import Image as PILImage
from pypdf import PdfReader, PdfWriter
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
app = FastAPI()
Session = object
get_db = lambda: None


def load_document_functions():
    source = (ROOT / 'main.py').read_text(encoding='utf-8-sig')
    start = source.index('DOCUMENT_TOOL_MAX_BYTES =')
    end = source.index('# Ephemeral WebRTC', start)
    tree = ast.parse(source[start:end])
    ast.increment_lineno(tree, source[:start].count('\n'))
    namespace = dict(globals())
    exec(compile(tree, str(ROOT / 'main.py'), 'exec'), namespace)
    return namespace


api = load_document_functions()


def pdf_bytes(pages=3):
    doc = pymupdf.open()
    for number in range(pages):
        page = doc.new_page(width=300, height=400)
        page.insert_text((20, 40), f'Page {number + 1}')
    result = doc.tobytes()
    doc.close()
    return result


async def response_bytes(response):
    return b''.join([chunk async for chunk in response.body_iterator])


class DocumentToolsTests(unittest.TestCase):
    def test_reorder_rotate_and_extract(self):
        async def run():
            response = await api['organize_pdf_pages'](
                [UploadFile(file=io.BytesIO(pdf_bytes()), filename='test.pdf')],
                json.dumps([[0, 2, 90], [0, 0, 0]]), 'test')
            reader = PdfReader(io.BytesIO(await response_bytes(response)))
            self.assertEqual([page.extract_text().strip() for page in reader.pages], ['Page 3', 'Page 1'])
            self.assertEqual(reader.pages[0].rotation, 90)
        asyncio.run(run())

    def test_organizer_accepts_200_pages_but_other_tools_keep_their_limit(self):
        data = pdf_bytes(200)
        self.assertEqual(len(api['_validated_pdf_reader'](data, max_pages=200).pages), 200)
        with self.assertRaises(HTTPException):
            api['_validated_pdf_reader'](data)

    def test_split_rejects_huge_ranges_without_expansion(self):
        started = time.monotonic()
        with self.assertRaises(HTTPException):
            api['_split_page_groups']('1-999999999999999999', 3)
        self.assertLess(time.monotonic() - started, 1)
        self.assertEqual(api['_split_page_groups']('1-2;3', 3), [[0, 1], [2]])
        for value in ('9' * 5000, '²', '0', '1-0'):
            with self.assertRaises(HTTPException):
                api['_split_page_groups'](value, 3)

    def test_bad_rotation_and_boolean_indices_are_rejected(self):
        for order in ([[0, 0, 45]], [[False, 0]]):
            with self.assertRaises(HTTPException):
                asyncio.run(api['organize_pdf_pages'](
                    [UploadFile(file=io.BytesIO(pdf_bytes()), filename='test.pdf')], json.dumps(order), 'test'))

    def test_pdf_word_export_is_explicitly_image_based(self):
        from docx import Document
        word = Document(io.BytesIO(api['_pdf_to_word_bytes'](pdf_bytes())))
        self.assertEqual(len(word.inline_shapes), 3)
        self.assertEqual(''.join(paragraph.text for paragraph in word.paragraphs).strip(), '')

    def test_ocr_word_round_trip_keeps_text_and_page_geometry(self):
        image = PILImage.new('RGB', (600, 800), 'white')
        artwork = io.BytesIO(); image.save(artwork, 'PNG')
        pages = []
        for text in ('Biology 12345', 'Second page'):
            pages.append({'width_px':600,'height_px':800,'width_pt':300,'height_pt':400,
                          'lines':[{'text':text,'x':40,'y':60,'width':250,'height':25}]})
        word = api['_ocr_layout_to_word_bytes']({'pages': pages}, [artwork.getvalue()] * 2)
        with patch.dict(api, {'_office_to_pdf_bytes': lambda *args: None}):
            output = api['_word_to_pdf_bytes'](word)
        pdf = pymupdf.open(stream=output, filetype='pdf')
        self.assertEqual(len(pdf), 2)
        self.assertIn('Biology 12345', pdf[0].get_text())
        self.assertIn('Second page', pdf[1].get_text())
        self.assertEqual(tuple(pdf[0].rect), (0, 0, 300, 400))
        pdf.close()

    def test_normal_word_never_silently_loses_layout_without_renderer(self):
        from docx import Document
        word = Document(); word.add_paragraph('Keep my content')
        data = io.BytesIO(); word.save(data)
        with patch.dict(api, {'_office_to_pdf_bytes': lambda *args: None}):
            with self.assertRaises(HTTPException) as caught:
                api['_word_to_pdf_bytes'](data.getvalue())
        self.assertEqual(caught.exception.status_code, 503)

    def test_office_renderer_has_isolated_profile_and_timeout(self):
        def render(args, **kwargs):
            output = Path(args[args.index('--outdir') + 1]) / 'document.pdf'
            output.write_bytes(pdf_bytes(1))
            self.assertTrue(any(arg.startswith('-env:UserInstallation=file:') for arg in args))
            self.assertEqual(kwargs['timeout'], 90)
        with patch.dict(os.environ, {'LIBREOFFICE_PATH':'test-office'}), patch.object(subprocess, 'run', side_effect=render):
            self.assertTrue(api['_office_to_pdf_bytes'](b'test', 'docx').startswith(b'%PDF'))


if __name__ == '__main__':
    unittest.main()
